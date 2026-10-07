# excel_export.py


import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from dados import normalizar

ESTILO_TITULO = Font(size=14, bold=True)
ESTILO_CABECALHO = Font(bold=True, color="FFFFFF")
PREENCHIMENTO_CABECALHO = PatternFill("solid", fgColor="4472C4")
ESTILO_TOTAL = Font(bold=True)
PREENCHIMENTO_TOTAL = PatternFill("solid", fgColor="D9E1F2")
BORDA_TOPO = Border(top=Side(style="thin"))
FORMATO_MOEDA = "R$ #,##0.00"


def _escrever_tabela(ws, df, linha_inicio, destacar_ultima_linha=False):
    colunas = list(df.columns)
    for j, nome_col in enumerate(colunas, start=1):
        cel = ws.cell(row=linha_inicio, column=j, value=nome_col)
        cel.font = ESTILO_CABECALHO
        cel.fill = PREENCHIMENTO_CABECALHO

    for i, (_, row) in enumerate(df.iterrows(), start=1):
        linha_atual = linha_inicio + i
        for j, nome_col in enumerate(colunas, start=1):
            valor = row[nome_col]
            if valor is not None and hasattr(valor, "item"):
                try:
                    valor = valor.item()
                except Exception:
                    pass
            cel = ws.cell(row=linha_atual, column=j, value=valor)
            if nome_col in ("VALOR_PREVISTO", "VALOR_EXECUTADO") and isinstance(valor, (int, float)):
                cel.number_format = FORMATO_MOEDA
            if destacar_ultima_linha and i == len(df):
                cel.font = ESTILO_TOTAL
                cel.fill = PREENCHIMENTO_TOTAL
                cel.border = BORDA_TOPO

    for j, nome_col in enumerate(colunas, start=1):
        maior = max([len(str(nome_col))] + [len(str(v)) for v in df[nome_col].astype(str)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(maior + 2, 10), 45)


def salvar_excel(municipio, dados_mun, painel_final, pasta_saida):
    wb = Workbook()

    ws_painel = wb.active
    ws_painel.title = "PAINEL_RESUMO"
    ws_painel["A1"] = f"PAINEL SETRAND - {municipio} (gerado em {datetime.now().strftime('%d/%m/%Y')})"
    ws_painel["A1"].font = ESTILO_TITULO
    _escrever_tabela(ws_painel, painel_final, linha_inicio=3, destacar_ultima_linha=True)

    ws_det = wb.create_sheet("DETALHAMENTO")
    _escrever_tabela(ws_det, dados_mun, linha_inicio=1)
    ws_det.freeze_panes = "A2"

    nome_arquivo = f"PAINEL_{normalizar(municipio).replace(' ', '_')}.xlsx"
    caminho_saida = os.path.join(pasta_saida, nome_arquivo)
    wb.save(caminho_saida)
    return caminho_saida
